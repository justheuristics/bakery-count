#!/usr/bin/env python3
"""
Monthly Code 206 price import — extracts that month's "Bakery Inventory Adjustment"
workbooks (confirmed prices for the month) into docs/price_import_<YYYY-MM>.json, which
app.js fetches at runtime (see computePriceImportPreview()).

Same rationale as the original August extractor: the write this feeds runs against
PRODUCTION, so the parse is done once, offline, self-verified, and committed as a plain
reviewable diff instead of parsing xlsx in the browser.

Inputs - a zip (or a folder) holding the month's workbooks:
  * the MAIN list: the workbook whose file name contains "FBK"  -> every store   (key "FBK")
  * any number of STORE-SPECIFIC lists: every other .xlsx       -> keyed "S<first store no.>"
    (e.g. "...2026 2_17_57.xlsx" -> S2, "...2026 22_350_121....xlsx" -> S22)

Rules baked in here (decided with the project owner, 2 Oct 2026):
  * Only the EX VAT column is imported. IN VAT is read for a consistency check only.
  * Main FBK price wins whenever a code is priced in FBK AND a store file.
  * A store-file price is used only for a code that FBK does not price at all (and only if
    every store file that prices it agrees - otherwise it is a CONFLICT and is skipped).
  * Rows whose price cell is "-" ("use store retail cost") carry no price and are never
    imported as one; they are listed under storeCostOnly for the human.
  * The app has ONE global price per item, so store membership of a file is context only.

Safety checks (refuses to write if any fails):
  * the Thai month + Buddhist year printed on each sheet ("ราคายืนยันระหว่าง วันที่ 1-30
    <เดือน> <พ.ศ.>") must match --month, so last month's files can't be imported as this
    month's (override only with --skip-month-check, and say why);
  * a code repeated inside one file must repeat the same price;
  * row accounting must add up; every file must have been recognised.
  For a month with hand-verified counts (KNOWN_EXPECTED) the per-file counts must match too.
  For any other month there is nothing to compare against - read the printed summary, it is
  the review. The JSON's own `summary` is re-verified by the app at run time either way.

Usage:
  python tools/extract_price_import_monthly.py --month 2026-10 ["/path/to/Bakery - October.zip"]
Writes: docs/price_import_<month>.json
"""
import argparse
import hashlib
import io
import json
import re
import sys
import zipfile
from pathlib import Path

import openpyxl

REPO_ROOT = Path(__file__).resolve().parent.parent
PRIORITY_MAIN = "FBK"

# Hand-verified against the actual workbooks. If the parser disagrees the parser is wrong
# (or the files changed) - tell the human, do not adjust the numbers.
KNOWN_EXPECTED = {
    "2026-09": {
        "FBK": {"table_rows": 242, "unique_codes": 240, "priced_unique": 240},
        "S22": {"table_rows": 83, "unique_codes": 82, "priced_unique": 54},
        "S2": {"table_rows": 61, "unique_codes": 59, "priced_unique": 34},
        "unique_codes_all_files": 312,
    },
}

THAI_MONTHS = ["มกราคม", "กุมภาพันธ์", "มีนาคม", "เมษายน", "พฤษภาคม", "มิถุนายน",
               "กรกฎาคม", "สิงหาคม", "กันยายน", "ตุลาคม", "พฤศจิกายน", "ธันวาคม"]


def num(x):
    return isinstance(x, (int, float)) and not isinstance(x, bool)


def store_key(member_name, taken):
    """FBK -> 'FBK'; otherwise 'S<first store no.>' from the 'a_b_c' store list in the file name."""
    if "FBK" in member_name:
        return PRIORITY_MAIN
    m = re.search(r"(\d+)(?:_\d+)+", member_name)
    key = "S" + m.group(1) if m else "S" + str(len(taken) + 1)
    base, i = key, 2
    while key in taken:
        key, i = f"{base}_{i}", i + 1
    return key


def parse_sheet(ws):
    """Return (header_lines, rows). Columns are located from the 'No.' header cell, because
    the files are laid out with different left margins."""
    grid = list(ws.iter_rows(values_only=True))
    hdr_idx = off = None
    for i, r in enumerate(grid):
        for j, v in enumerate(r):
            if isinstance(v, str) and v.strip() == "No.":
                hdr_idx, off = i, j
                break
        if hdr_idx is not None:
            break
    if hdr_idx is None:
        raise SystemExit("ERROR: no 'No.' header row found in sheet " + ws.title)
    header_lines = [str(v).strip() for r in grid[:hdr_idx] for v in r if isinstance(v, str) and v.strip()]
    rows = []
    for i, r in enumerate(grid[hdr_idx + 1:], start=hdr_idx + 2):
        no = r[off]
        code = r[off + 3] if len(r) > off + 3 else None
        if not num(no) or code in (None, ""):
            continue  # EX/IN VAT sub-header row, signature block, blanks
        ex, inv = r[off + 5], r[off + 6]
        rows.append({
            "sheetRow": i,
            "code": str(int(code)) if num(code) else str(code).strip(),
            "name": str(r[off + 4]).strip() if r[off + 4] is not None else "",
            "class": str(r[off + 1]).strip() if r[off + 1] is not None else "",
            "buyer": str(r[off + 2]).strip() if r[off + 2] is not None else "",
            "exVat": float(ex) if num(ex) else None,
            "inVat": float(inv) if num(inv) else None,
            "exRaw": None if num(ex) else (str(ex).strip() if ex is not None else None),
        })
    return header_lines, rows


def sheet_month(header_lines):
    """The 'YYYY-MM' the sheet says its prices are confirmed for, or None if not found."""
    for line in header_lines:
        if "ราคายืนยัน" not in line:
            continue
        for idx, name in enumerate(THAI_MONTHS, start=1):
            m = re.search(re.escape(name) + r"\s*(\d{4})", line)
            if m:
                return f"{int(m.group(1)) - 543:04d}-{idx:02d}"
    return None


def load_workbooks(source):
    """Yield (member_name, bytes) for every .xlsx in a zip file or a folder."""
    p = Path(source)
    if p.is_dir():
        for f in sorted(p.glob("*.xlsx")):
            if not f.name.startswith("~$"):
                yield f.name, f.read_bytes()
        return
    with zipfile.ZipFile(p) as z:
        for n in z.namelist():
            if n.lower().endswith(".xlsx") and not n.startswith("__MACOSX") and not Path(n).name.startswith("~$"):
                yield n, z.read(n)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--month", required=True, help="YYYY-MM the prices are confirmed for, e.g. 2026-10")
    ap.add_argument("source", nargs="?", help="zip or folder with the month's workbooks (default: ~/Downloads/Bakery - <Month>.zip is NOT guessed - pass it)")
    ap.add_argument("--skip-month-check", action="store_true", help="do not require the sheets' printed month to equal --month")
    args = ap.parse_args()
    if not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", args.month):
        sys.exit("ERROR: --month must look like 2026-10")
    month = args.month
    if not args.source:
        sys.exit("ERROR: pass the zip (or folder) holding the month's workbooks")
    if not Path(args.source).exists():
        sys.exit(f"ERROR: not found: {args.source}")
    json_path = REPO_ROOT / "docs" / f"price_import_{month}.json"

    files, per_file, header_notes, sheet_months = {}, {}, {}, {}
    for member, data in load_workbooks(args.source):
        key = store_key(member, files)
        files[key] = {"file": Path(member).name, "sha256": hashlib.sha256(data).hexdigest()}
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        hl, rows = parse_sheet(wb.worksheets[0])
        per_file[key] = rows
        header_notes[key] = [t for t in hl if "สาขา" in t or "ราคายืนยัน" in t]
        sheet_months[key] = sheet_month(hl)
    if PRIORITY_MAIN not in per_file:
        sys.exit('ERROR: no main workbook found (a file name containing "FBK")')
    PRIORITY = [PRIORITY_MAIN] + sorted(k for k in per_file if k != PRIORITY_MAIN)

    problems = []

    # ── the sheets must say they are for the month we were asked to import ──
    for k in PRIORITY:
        if sheet_months[k] is None:
            msg = f"{k} ({files[k]['file']}): could not read the confirmed-price month from the sheet header"
            (print("WARNING:", msg) if args.skip_month_check else problems.append(msg))
        elif sheet_months[k] != month and not args.skip_month_check:
            problems.append(f"{k} ({files[k]['file']}): sheet says prices are for {sheet_months[k]}, not {month}")

    # ── per-file de-dup: a repeated code must repeat the same price, else refuse ──
    uniq = {}
    for key, rows in per_file.items():
        seen = {}
        for r in rows:
            prev = seen.get(r["code"])
            if prev is None:
                seen[r["code"]] = r
            elif prev["exVat"] != r["exVat"]:
                problems.append(f"{key}: code {r['code']} repeated with different prices {prev['exVat']} vs {r['exVat']}")
        uniq[key] = seen

    counts = {k: {"table_rows": len(per_file[k]), "unique_codes": len(uniq[k]),
                  "priced_unique": sum(1 for r in uniq[k].values() if r["exVat"] is not None)} for k in PRIORITY}
    all_codes = set().union(*[set(u) for u in uniq.values()])
    counts["unique_codes_all_files"] = len(all_codes)
    known = KNOWN_EXPECTED.get(month)
    if known:
        if set(known) - {"unique_codes_all_files"} != set(PRIORITY):
            problems.append(f"files found {sorted(PRIORITY)} differ from the hand-verified set {sorted(set(known) - {'unique_codes_all_files'})}")
        for k, want in known.items():
            if k == "unique_codes_all_files":
                if counts[k] != want:
                    problems.append(f"unique codes across files: {counts[k]} != {want}")
            elif k in counts:
                for f, w in want.items():
                    if counts[k][f] != w:
                        problems.append(f"{k}.{f}: {counts[k][f]} != {w}")

    # ── VAT consistency: IN should be EX x 1.07. Flag (never silently 'fix') the rest ──
    vat_flags = []
    for key in PRIORITY:
        for r in uniq[key].values():
            if r["exVat"] and r["inVat"] and abs(r["inVat"] / r["exVat"] - 1.07) > 0.002:
                vat_flags.append({"file": key, "code": r["code"], "name": r["name"], "exVat": r["exVat"], "inVat": r["inVat"],
                                  "note": "IN VAT column is not EX x 1.07 - EX VAT column is what is imported"})

    # ── resolve one row per code ──
    out_rows, store_cost_only, conflicts, fbk_overrides_store_dash = [], [], [], []
    for code in sorted(all_codes, key=lambda c: (len(c), c)):
        by_file = {k: uniq[k][code] for k in PRIORITY if code in uniq[k]}
        priced = {k: r for k, r in by_file.items() if r["exVat"] is not None}
        any_row = next(iter(by_file.values()))
        name = by_file["FBK"]["name"] if "FBK" in by_file else any_row["name"]
        if not priced:
            store_cost_only.append({"code": code, "name": name, "files": sorted(by_file)})
            continue
        if "FBK" in priced:
            src = "FBK"
            for k, r in by_file.items():
                if k != "FBK" and r["exVat"] is None:
                    fbk_overrides_store_dash.append({"code": code, "name": name, "file": k, "fbkExVat": priced["FBK"]["exVat"],
                                                     "note": "store file says 'use store retail cost'; main FBK price used"})
        else:
            vals = {round(r["exVat"], 4) for r in priced.values()}
            if len(vals) > 1:
                conflicts.append({"code": code, "name": name, "kind": "STORE_FILES_DISAGREE",
                                  "prices": {k: r["exVat"] for k, r in priced.items()}})
                out_rows.append({"code": code, "name": name, "class": any_row["class"], "exVat": None, "inVat": None,
                                 "source": None, "status": "CONFLICT_UNRESOLVED", "sheetsSeen": sorted(by_file)})
                continue
            src = next(k for k in PRIORITY if k in priced)
        row = priced[src]
        others = {k: r["exVat"] for k, r in priced.items() if k != src and abs(r["exVat"] - row["exVat"]) > 0.0049}
        if others:
            conflicts.append({"code": code, "name": name, "kind": "FBK_WINS" if src == "FBK" else "STORE_FILE",
                              "used": {src: row["exVat"]}, "overruled": others})
        out_rows.append({"code": code, "name": name, "class": row["class"], "exVat": row["exVat"], "inVat": row["inVat"],
                         "source": src, "status": "PRICED", "sheetsSeen": sorted(by_file)})

    n_priced = sum(1 for r in out_rows if r["status"] == "PRICED")
    summary = {
        "unique_codes_all_files": len(all_codes),
        "priced_rows": n_priced,
        "from_FBK": sum(1 for r in out_rows if r["source"] == "FBK"),
        "from_store_file_only": sum(1 for r in out_rows if r["status"] == "PRICED" and r["source"] != "FBK"),
        "conflict_unresolved": sum(1 for r in out_rows if r["status"] == "CONFLICT_UNRESOLVED"),
        "store_cost_only": len(store_cost_only),
    }
    if summary["priced_rows"] + summary["conflict_unresolved"] + summary["store_cost_only"] != len(all_codes):
        problems.append("row accounting does not add up to the unique code count")

    if problems:
        print("REFUSING TO WRITE - self-checks failed:", file=sys.stderr)
        for p in problems:
            print("  -", p, file=sys.stderr)
        sys.exit(1)

    doc = {
        "effectiveFrom": month,
        "basis": "EX_VAT",
        "priceSource": f"code-206-{month}",
        "generatedBy": "tools/extract_price_import_monthly.py",
        "sourceFiles": files,
        "sourceHeaderNotes": header_notes,
        "expected": known or counts,
        "summary": summary,
        "rows": out_rows,
        "storeCostOnly": store_cost_only,
        "conflicts": conflicts,
        "fbkOverridesStoreDash": fbk_overrides_store_dash,
        "vatInconsistent": vat_flags,
    }
    json_path.write_text(json.dumps(doc, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print("wrote", json_path)
    print("  files:", {k: (files[k]["file"][-40:], sheet_months[k]) for k in PRIORITY})
    print("  per-file counts:", counts)
    print("  summary:", summary)
    print("  conflicts:", len(conflicts), "| fbk-over-store-dash:", len(fbk_overrides_store_dash), "| vat flags:", [v["code"] for v in vat_flags])
    if not known:
        print(f"  NOTE: no hand-verified counts exist for {month} - review the per-file counts above against the workbooks.")


if __name__ == "__main__":
    main()
